import unittest
from io import BytesIO
from openpyxl import Workbook
from import_world_bank import parse, number, DATASET, ENERGY, METALS

class PinkSheetTests(unittest.TestCase):
    def workbook(self):
        w=Workbook();s=w.active;s.title='Monthly Prices';d=w.create_sheet('Description')
        s.cell(1,1,DATASET);s.cell(4,1,'Updated on October 02, 2026')
        for col,(header,label,row,unit) in enumerate(ENERGY+METALS,2):
            s.cell(5,col,header);s.cell(6,col,unit);s.cell(7,col,10);s.cell(8,col,'..')
            d.cell(row,2,label);d.cell(row,20,'Bloomberg Finance L.P.; World Bank.' if (header,label,row,unit) in METALS else 'World Bank.')
        s.cell(7,1,'2026M08');s.cell(8,1,'2026M09')
        return w
    def parse(self,w):
        b=BytesIO();w.save(b);return parse(b.getvalue())
    def test_nulls_never_zero(self):
        for v in [None,'..','…','W','N/D']:self.assertIsNone(number(v))
        with self.assertRaises(ValueError):number('bogus')
    def test_metadata_and_restricted_provider_preserved(self):
        date,energy,metals=self.parse(self.workbook())
        self.assertEqual(date,'2026-10-02');self.assertEqual(len(metals),10)
        self.assertEqual(metals[0]['rows'][0],{'label':'2026-09','values':[None]})
        self.assertEqual(metals[0]['rows'][1]['values'],[10])
        self.assertIn('Bloomberg Finance L.P.',metals[0]['attribution'])
        self.assertIn('unverified',metals[0]['licenseStatus'])
        self.assertTrue(all(r['values']==[None] for r in metals[2]['rows']))
        self.assertEqual(energy[0]['rows'][1]['values'],[10])
    def test_unit_or_provider_changes_stop_import(self):
        w=self.workbook();w['Monthly Prices'].cell(6,2,'EUR/t')
        with self.assertRaises(ValueError):self.parse(w)
        w=self.workbook();w['Description'].cell(88,20).value=None
        with self.assertRaises(ValueError):self.parse(w)

if __name__=='__main__':unittest.main()
