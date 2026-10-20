"""Manual Pink Sheet snapshot. Review source/units/rights and diff before publishing."""
import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import re
from urllib.request import urlopen
from openpyxl import load_workbook

URL = 'https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx'
TERMS = 'https://www.worldbank.org/en/about/legal/terms-of-use-for-datasets'
DATASET = 'World Bank Commodity Price Data (The Pink Sheet)'
# Workbook header, Spanish label, Description row, expected source unit.
ENERGY = [('Crude oil, average', 'Petróleo · media Brent, Dubái y WTI', 7, '($/bbl)'),
          ('Natural gas index', 'Gas natural · índice', 11, '(2010=100)')]
METALS = [('Aluminum', 'Aluminio', 88, '($/mt)'), ('Copper', 'Cobre', 89, '($/mt)'),
          ('Iron ore, cfr spot', 'Mineral de hierro', 91, '($/dmtu)'),
          ('Lead', 'Plomo', 92, '($/mt)'), ('Nickel', 'Níquel', 93, '($/mt)'),
          ('Tin', 'Estaño', 99, '($/mt)'), ('Zinc', 'Zinc', 100, '($/mt)'),
          ('Gold', 'Oro', 103, '($/troy oz)'), ('Platinum', 'Platino', 104, '($/troy oz)'),
          ('Silver', 'Plata', 105, '($/troy oz)')]


def number(v):
    if v is None or str(v).strip() in {'', '…', '...', '..', 'W', 'N/D', 'NA'}:
        return None
    if isinstance(v, bool) or not isinstance(v, (float, int)) or not math.isfinite(v) or v < 0:
        raise ValueError(f'Unexpected source value: {v!r}')
    return v


def parse(raw):
    w = load_workbook(BytesIO(raw), data_only=True)
    s, d = w['Monthly Prices'], w['Description']
    if s.cell(1, 1).value != DATASET:
        raise ValueError('Unexpected dataset')
    updated = datetime.strptime(s.cell(4, 1).value, 'Updated on %B %d, %Y').date().isoformat()
    periods = [r for r in s.iter_rows(min_row=7, values_only=True) if re.fullmatch(r'\d{4}M\d{2}', str(r[0]))]
    if not periods or [r[0] for r in periods] != sorted(set(r[0] for r in periods)):
        raise ValueError('Missing, duplicate or unordered months')
    headers = {c.value: c.column for c in s[5] if c.value}
    def series(spec):
        header, label, description_row, expected_unit = spec
        col = headers[header]
        if s.cell(6, col).value != expected_unit:
            raise ValueError('Source unit changed')
        providers = d.cell(description_row, 20).value
        if not providers or not str(d.cell(description_row, 2).value).strip():
            raise ValueError('Missing provider metadata')
        unit = {'($/mt)': 'USD/t', '($/bbl)': 'USD/barril', '(2010=100)': 'índice 2010=100',
                '($/dmtu)': 'USD/dmtu (unidad de la hoja)', '($/troy oz)': 'USD/oz troy'}[expected_unit]
        # Iron header says dmtu while Description says dollar/dry ton: do not guess.
        ambiguous_unit = header == 'Iron ore, cfr spot'
        rows = [{'label': r[0].replace('M', '-'), 'values': [None if ambiguous_unit else number(r[col-1])]} for r in periods[-24:]]
        return {'title': f'{label} · precio mensual' if header != 'Natural gas index' else f'{label} mensual',
                'columns': ['Media mensual'], 'unit': unit, 'rows': list(reversed(rows)), 'source': URL,
                'sourceName': 'World Bank · Pink Sheet', 'vintage': updated, 'published': updated,
                'kind': 'reportado', 'monthly': True,
                'attribution': f'The World Bank: {DATASET}: {providers}', 'terms': TERMS,
                'providers': providers, 'sourceHeader': header, 'sourceUnit': expected_unit,
                'sourceDescription': d.cell(description_row, 2).value,
                'licenseStatus': 'third_party_rights_unverified_owner_requested_republication' if spec in METALS else 'CC BY 4.0 + World Bank dataset terms',
                **({'unavailableReason': 'N/D · unidad no resuelta: cabecera $/dmtu; descripción USD/t seca.'} if ambiguous_unit else {})}
    return updated, [series(x) for x in ENERGY], [series(x) for x in METALS]


def main():
    p = argparse.ArgumentParser(); p.add_argument('--input'); args = p.parse_args()
    raw = Path(args.input).read_bytes() if args.input else urlopen(URL, timeout=60).read()
    updated, energy, metals = parse(raw)
    root = Path(__file__).resolve().parents[2] / 'data/theme-fundamentals'
    for name, title, series in [('utilities', 'Energía y utilities', energy), ('materiales', 'Materiales críticos', metals)]:
        file = root / f'{name}.json'
        data = json.loads(file.read_text()) if file.exists() else {'id': name, 'title': title, 'metrics': [], 'series': [], 'notes': [], 'links': []}
        data['checkedAt'] = datetime.now(timezone.utc).date().isoformat()
        data['series'] = series + [s for s in data['series'] if not s.get('monthly')]
        data['notes'] = [n for n in data['notes'] if not n.startswith('Precios mensuales externos:')]
        data['pinkSheet'] = {'source': URL, 'sha256': hashlib.sha256(raw).hexdigest(), 'vintage': updated,
                            'window': 'Últimos 24 meses disponibles en el XLS; medias mensuales nominales, no precios en vivo.',
                            'terms': TERMS}
        file.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')

if __name__ == '__main__':
    main()
