"""Datasets gratuitos de Aswath Damodaran (NYU Stern) para el WACC.

Conexion 10c del mapa de conexiones. Damodaran publica cada mes la prima de
riesgo implicita del S&P 500 («Implied ERP by month», ERPbymonth.xlsx). El
WACC usaba una prima fija de settings (0,0433 = enero de 2025 en esa misma
serie, columna «ERP (T12m)»); con este conector se usa la del mes vigente.

Fuente: Aswath Damodaran, «Implied Equity Risk Premiums», NYU Stern,
https://pages.stern.nyu.edu/~adamodar/New_Home_Page/datafile/implpr.html
(datos gratuitos; citar la fuente). Se usa la columna «ERP (T12 m with
sustainable payout)»: en el backtest 2010–2026 dio mejor IC agregado y mejor
calibración que «ERP (T12m)», que además está en mínimo histórico (payout de
recompras no sostenible). La T12m se guarda solo como referencia en los
metadatos. El valor de settings (0,0433) sigue siendo el respaldo sin red.

Fail-closed: sin red, fichero cambiado o fila invalida -> lista vacia / None,
y el WACC vuelve al valor de settings declarando la procedencia. Una descarga
por proceso cada ``CACHE_TTL_SECONDS`` (el fichero cambia una vez al mes).
"""

from __future__ import annotations

import io
import logging
import time
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import httpx

from app.services.data_source_health import get_data_source_registry

# Clave en el registry de salud de fuentes. La URL es fija (un xlsx mensual):
# cualquier fallo de descarga o un xlsx sin filas validas es de la fuente.
REGISTRY_KEY = "damodaran"

logger = logging.getLogger("gerarai.damodaran")

DAMODARAN_ERP_URL = "https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx"
DAMODARAN_ERP_PAGE = "https://pages.stern.nyu.edu/~adamodar/New_Home_Page/datafile/implpr.html"
DAMODARAN_ERP_SHEET = "Historical ERP"
DAMODARAN_ERP_COLUMN = "ERP (T12 m with sustainable payout)"
_T12M_REFERENCE_COLUMN = "ERP (T12m)"
_TBOND_COLUMN = "T.Bond Rate"
_HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; GerarAI/0.1; +mailto:contact@example.com)"}
CACHE_TTL_SECONDS = 12 * 3600
# Una ERP del S&P 500 fuera de (0, 15 %) no es un dato: es una celda movida.
_MAX_PLAUSIBLE_ERP = Decimal("0.15")


@dataclass(frozen=True)
class ImpliedErp:
    month: date
    erp: Decimal
    erp_t12m: Decimal | None
    tbond: Decimal | None


_cache: dict[str, tuple[float, list[ImpliedErp]]] = {}


def reset_cache() -> None:
    _cache.clear()


def _decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def parse_implied_erp_xlsx(content: bytes) -> list[ImpliedErp]:
    """Filas mensuales validas de la hoja «Historical ERP», en orden cronologico."""
    try:
        import openpyxl

        with warnings.catch_warnings():
            # El xlsx de Damodaran trae extensiones de Excel que openpyxl no
            # conoce («Unknown extension»): ruido, no afecta a los datos.
            warnings.simplefilter("ignore", UserWarning)
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - fichero corrupto o cambiado: fail-closed
        logger.warning("Damodaran ERP: no se pudo abrir el xlsx", exc_info=True)
        return []
    if DAMODARAN_ERP_SHEET not in workbook.sheetnames:
        return []
    rows = workbook[DAMODARAN_ERP_SHEET].iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    names = [str(cell).strip() if cell is not None else "" for cell in header]
    if DAMODARAN_ERP_COLUMN not in names or not names[0].lower().startswith("start of month"):
        return []
    erp_index = names.index(DAMODARAN_ERP_COLUMN)
    t12m_index = names.index(_T12M_REFERENCE_COLUMN) if _T12M_REFERENCE_COLUMN in names else None
    tbond_index = names.index(_TBOND_COLUMN) if _TBOND_COLUMN in names else None
    out: dict[date, ImpliedErp] = {}
    for row in rows:
        if not row or not isinstance(row[0], datetime | date):
            continue
        erp = _decimal(row[erp_index]) if erp_index < len(row) else None
        if erp is None or not (Decimal("0") < erp < _MAX_PLAUSIBLE_ERP):
            continue
        month = row[0].date() if isinstance(row[0], datetime) else row[0]
        out[month] = ImpliedErp(
            month=month,
            erp=erp,
            erp_t12m=_decimal(row[t12m_index]) if t12m_index is not None and t12m_index < len(row) else None,
            tbond=_decimal(row[tbond_index]) if tbond_index is not None and tbond_index < len(row) else None,
        )
    return [out[month] for month in sorted(out)]


def latest_implied_erp(rows: list[ImpliedErp], as_of: date) -> ImpliedErp | None:
    """Ultimo mes publicado en o antes de ``as_of`` (nunca un mes futuro)."""
    eligible = [row for row in rows if row.month <= as_of]
    return eligible[-1] if eligible else None


async def _download(url: str) -> bytes:
    async with httpx.AsyncClient(timeout=20.0, headers=_HEADERS, follow_redirects=True) as client:
        response = await client.get(url)
        response.raise_for_status()
        return response.content


async def implied_erp_series() -> list[ImpliedErp]:
    """Serie mensual de ERP implicito (cacheada en proceso). ``[]`` si no hay dato."""
    cached = _cache.get(DAMODARAN_ERP_URL)
    if cached and time.monotonic() - cached[0] < CACHE_TTL_SECONDS:
        return cached[1]
    registry = get_data_source_registry()
    started = time.monotonic()
    try:
        content = await _download(DAMODARAN_ERP_URL)
    except Exception as exc:  # noqa: BLE001 - red caida o 4xx/5xx: el WACC usa settings
        registry.record_failure(REGISTRY_KEY, f"descarga fallida: {type(exc).__name__}")
        logger.warning("Damodaran ERP: descarga fallida", exc_info=True)
        return []
    rows = parse_implied_erp_xlsx(content)
    if rows:
        _cache[DAMODARAN_ERP_URL] = (time.monotonic(), rows)
        registry.record_success(REGISTRY_KEY, (time.monotonic() - started) * 1000)
    else:
        # Descarga OK pero sin filas validas: el fichero cambio de formato.
        registry.record_failure(REGISTRY_KEY, "xlsx sin filas validas (formato cambiado)")
    return rows
