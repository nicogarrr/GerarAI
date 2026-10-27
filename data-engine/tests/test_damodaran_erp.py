"""Damodaran (NYU Stern): ERP implicito mensual como input del WACC (conexion 10c).

Fuente gratuita y citada: «Implied ERP by month» de Aswath Damodaran
(ERPbymonth.xlsx). El ERP fijo de settings (4,33 %) es el de enero de 2025 en
esa misma serie (columna «ERP (T12m)»); con el dato mensual el WACC deja de
usar una prima congelada. Sin red o sin dato, se vuelve al valor de settings
y se declara la procedencia.
"""

import asyncio
import io
from datetime import date, datetime
from decimal import Decimal

import httpx
import openpyxl
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.models.entities import Base, Company, Document, FinancialFact
from app.services import wacc_input_service
from app.services.connectors import damodaran
from app.services.connectors.damodaran import (
    DAMODARAN_ERP_URL,
    ImpliedErp,
    latest_implied_erp,
    parse_implied_erp_xlsx,
)
from app.services.wacc_input_service import WaccInputService


def _workbook(rows: list[tuple]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Historical ERP"
    ws.append((
        "Start of month", "S&P 500", "T.Bond Rate", "$ Riskfree Rate", "Ten-year average CF",
        "CF (Trailing 12 month)", "Normalized CF", "Expected growth rate",
        "ERP (T12 m with sustainable payout)", "ERP (T12m)",
    ))
    for row in rows:
        ws.append(row)
    other = wb.create_sheet("Last 12 months data")
    other.append(("Date", "S&P 500", "10-year US Treasury", "ERP"))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


SAMPLE = _workbook([
    (datetime(2025, 1, 1), 5881, 0.0458, None, None, None, None, None, 0.04, 0.0433),
    (datetime(2026, 9, 1), 7686.14, 0.0475, 0.0453, 301.16, 201.83, 175.27, 0.1399, 0.0414, 0.0409),
    (datetime(2026, 10, 1), 7652, 0.0529, 0.0507, 306.5, 198.17, 176.54, 0.1459, 0.042, 0.037),
    (None, None, None, None, None, None, None, None, None, None),
    ("Nota: fuente Damodaran", None, None, None, None, None, None, None, None, None),
    (datetime(2026, 11, 1), 7700, 0.05, None, None, None, None, None, None, "n/a"),
    (datetime(2026, 12, 1), 7700, 0.05, None, None, None, None, None, None, 0.9),
])


def test_parse_lee_la_columna_erp_t12m_y_descarta_filas_invalidas():
    rows = parse_implied_erp_xlsx(SAMPLE)
    assert [r.month for r in rows] == [date(2025, 1, 1), date(2026, 9, 1), date(2026, 10, 1)]
    october = rows[-1]
    assert october.erp == Decimal("0.037")
    assert october.erp_sustainable_payout == Decimal("0.042")
    assert october.tbond == Decimal("0.0529")
    # El ERP fijo de settings (0,0433) es el de enero de 2025 en esta serie.
    assert rows[0].erp == Decimal("0.0433")


def test_parse_sin_hoja_o_cabecera_esperada_falla_cerrado():
    wb = openpyxl.Workbook()
    buffer = io.BytesIO()
    wb.save(buffer)
    assert parse_implied_erp_xlsx(buffer.getvalue()) == []
    assert parse_implied_erp_xlsx(b"no es un xlsx") == []


def test_latest_implied_erp_no_usa_meses_futuros():
    rows = parse_implied_erp_xlsx(SAMPLE)
    assert latest_implied_erp(rows, date(2026, 10, 27)).month == date(2026, 10, 1)
    assert latest_implied_erp(rows, date(2026, 9, 30)).month == date(2026, 9, 1)
    assert latest_implied_erp(rows, date(2024, 1, 1)) is None
    assert latest_implied_erp([], date(2026, 10, 27)) is None


def test_fetch_usa_cache_en_proceso(monkeypatch):
    calls: list[str] = []

    async def fake_download(url: str) -> bytes:
        calls.append(url)
        return SAMPLE

    damodaran.reset_cache()
    monkeypatch.setattr(damodaran, "_download", fake_download)
    first = asyncio.run(damodaran.implied_erp_series())
    second = asyncio.run(damodaran.implied_erp_series())
    assert first == second and len(first) == 3
    assert calls == [DAMODARAN_ERP_URL]
    damodaran.reset_cache()


def test_fetch_con_error_de_red_devuelve_lista_vacia(monkeypatch):
    async def boom(url: str) -> bytes:
        raise httpx.ConnectError("sin red")

    damodaran.reset_cache()
    monkeypatch.setattr(damodaran, "_download", boom)
    assert asyncio.run(damodaran.implied_erp_series()) == []
    damodaran.reset_cache()


# -- integracion con WaccInputService ------------------------------------


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.info["tenant_id"] = "tenant-test"
        yield session


class _StubFred:
    def configured(self) -> bool:
        return True

    async def series(self, series_id: str, limit: int = 10) -> dict:
        return {"observations": [{"date": "2026-09-21", "value": "4.25"}]}

    async def series_csv(self, series_id, limit=10):
        return {"observations": []}


def _setup(monkeypatch, erp: ImpliedErp | None):
    async def _empty(self, ticker):
        return {}

    async def _erp(self, as_of):
        return erp

    monkeypatch.setattr(WaccInputService, "_yfinance_market_inputs", _empty)
    monkeypatch.setattr(WaccInputService, "_damodaran_implied_erp", _erp)
    monkeypatch.setattr(wacc_input_service, "FREDClient", lambda *a, **k: _StubFred())


def _company(db: Session) -> Company:
    company = Company(
        ticker="UBER", name="Uber", exchange="NYSE", currency="USD",
        sector="Tech", industry="Tech", company_type="holding",
        valuation_model="unassigned", special_sources=[], special_risks=[], factor_tags=[],
    )
    db.add(company)
    db.commit()
    return company


def _erp_fact(db, company) -> FinancialFact:
    fact = db.scalar(
        select(FinancialFact).where(
            FinancialFact.company_id == company.id,
            FinancialFact.metric == "equity_risk_premium",
        )
    )
    assert fact is not None
    return fact


def _doc(db, company, source_type: str) -> Document | None:
    return db.scalar(
        select(Document).where(Document.company_id == company.id, Document.source_type == source_type)
    )


def test_refresh_usa_el_erp_implicito_de_damodaran_y_lo_cita(db, monkeypatch):
    _setup(monkeypatch, ImpliedErp(date(2026, 10, 1), Decimal("0.037"), Decimal("0.042"), Decimal("0.0529")))
    company = _company(db)
    asyncio.run(WaccInputService().refresh(db, company))

    fact = _erp_fact(db, company)
    assert fact.value == Decimal("0.037")
    assert fact.confidence == Decimal("0.80")
    source = db.get(Document, fact.source_id)
    assert source is not None and source.source_type == "damodaran"
    assert source.metadata_["source_url"] == DAMODARAN_ERP_URL
    assert source.metadata_["month"] == "2026-10-01"
    assert source.metadata_["column"] == "ERP (T12m)"
    assert source.metadata_["erp_sustainable_payout"] == "0.042"
    assert "Damodaran" in source.title
    policy = _doc(db, company, "wacc_policy")
    assert policy is not None
    assert policy.metadata_["equity_risk_premium_source"] == "damodaran_implied_erp"


def test_refresh_sin_dato_de_damodaran_usa_settings_y_lo_declara(db, monkeypatch):
    _setup(monkeypatch, None)
    company = _company(db)
    asyncio.run(WaccInputService().refresh(db, company))

    fact = _erp_fact(db, company)
    assert fact.value == Decimal("0.0433")
    assert fact.confidence == Decimal("0.70")
    assert _doc(db, company, "damodaran") is None
    policy = _doc(db, company, "wacc_policy")
    assert policy is not None
    assert policy.metadata_["equity_risk_premium_source"] == "settings_default"


def test_damodaran_implied_erp_respeta_el_interruptor(monkeypatch):
    async def must_not_fetch():
        raise AssertionError("con WACC_DAMODARAN_ERP=0 no se descarga nada")

    monkeypatch.setenv("WACC_DAMODARAN_ERP", "0")
    monkeypatch.setattr(wacc_input_service, "implied_erp_series", must_not_fetch)
    assert asyncio.run(WaccInputService()._damodaran_implied_erp("2026-10-27")) is None


def test_damodaran_implied_erp_elige_el_mes_vigente(monkeypatch):
    async def series():
        return parse_implied_erp_xlsx(SAMPLE)

    monkeypatch.delenv("WACC_DAMODARAN_ERP", raising=False)
    monkeypatch.setattr(wacc_input_service, "implied_erp_series", series)
    picked = asyncio.run(WaccInputService()._damodaran_implied_erp("2026-10-27"))
    assert picked is not None and picked.month == date(2026, 10, 1)
