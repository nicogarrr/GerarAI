"""Ingesta DOCX/XLSX: un ZIP muy comprimible (zip bomb) se rechaza antes de parsear."""

import zipfile
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.services import document_ingestion_service as ingestion
from app.services.document_ingestion_service import DocumentIngestionService, _guard_ooxml_zip


def _zip(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def _real_docx() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Margen operativo estable y guia de crecimiento del 12 %.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _real_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["metric", "value"])
    sheet.append(["revenue_growth", "12%"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_docx_with_extreme_compression_ratio_is_rejected_before_parsing():
    # 8 MB de ceros comprimen a unos pocos KB: ratio >> 200.
    bomb = _zip({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"\0" * (8 * 1024 * 1024)})
    assert len(bomb) < 100 * 1024
    with pytest.raises(ValueError, match="zip bomb"):
        DocumentIngestionService()._parse_native(bomb, ".docx")


def test_xlsx_over_total_uncompressed_limit_is_rejected(monkeypatch):
    monkeypatch.setattr(ingestion, "MAX_OOXML_UNCOMPRESSED_BYTES", 512 * 1024)
    payload = _zip({f"xl/worksheets/sheet{i}.xml": bytes(range(256)) * 1024 for i in range(4)})
    with pytest.raises(ValueError, match="supera el limite"):
        DocumentIngestionService()._parse_native(payload, ".xlsx")


def test_too_many_entries_is_rejected(monkeypatch):
    monkeypatch.setattr(ingestion, "MAX_OOXML_ENTRIES", 5)
    payload = _zip({f"part{i}.xml": b"<x/>" for i in range(6)})
    with pytest.raises(ValueError, match="demasiadas entradas"):
        _guard_ooxml_zip(payload)


def test_corrupt_zip_is_a_clean_value_error():
    with pytest.raises(ValueError, match="ZIP corrupto"):
        DocumentIngestionService()._parse_native(b"esto no es un zip", ".docx")


def test_real_docx_and_xlsx_still_parse():
    docx_parsed = DocumentIngestionService()._parse_native(_real_docx(), ".docx")
    assert docx_parsed.parser == "python_docx"
    assert "Margen operativo" in docx_parsed.blocks[0].text

    xlsx_parsed = DocumentIngestionService()._parse_native(_real_xlsx(), ".xlsx")
    assert xlsx_parsed.parser == "openpyxl"
    assert "revenue_growth" in xlsx_parsed.blocks[0].text


def _forge_central_size(payload: bytes, size: int) -> bytes:
    import struct

    result = bytearray(payload)
    offset = result.index(b"PK\x01\x02")
    struct.pack_into("<I", result, offset + 24, size)
    return bytes(result)


def test_forged_directory_size_cannot_hide_real_expansion(monkeypatch):
    # Acotado: 256 KiB reales, limite 128 KiB; no agota memoria.
    payload = _forge_central_size(_zip({"word/document.xml": b"x" * (256 * 1024)}), 1)
    monkeypatch.setattr(ingestion, "MAX_OOXML_UNCOMPRESSED_BYTES", 128 * 1024)
    monkeypatch.setattr(
        DocumentIngestionService, "_parse_docx",
        lambda *args: pytest.fail("el parser no debe recibir un ZIP que excede el limite"),
    )
    with pytest.raises(ValueError, match="limite real"):
        DocumentIngestionService()._parse_native(payload, ".docx")


def test_small_forged_size_is_rejected_even_below_real_limit():
    payload = _forge_central_size(_zip({"word/document.xml": b"abc" * 100}), 1)
    with pytest.raises(ValueError, match="inconsistentes"):
        _guard_ooxml_zip(payload)


def test_multiple_forged_entries_are_rejected_before_parser(monkeypatch):
    import struct

    payload = bytearray(_zip({"a.xml": b"x" * 80_000, "b.xml": b"y" * 80_000}))
    pos = 0
    for _ in range(2):
        pos = payload.index(b"PK\x01\x02", pos)
        struct.pack_into("<I", payload, pos + 24, 1)
        pos += 4
    monkeypatch.setattr(ingestion, "MAX_OOXML_UNCOMPRESSED_BYTES", 128_000)
    # Primera entrada se rechaza por discordancia antes de llegar al parser.
    with pytest.raises(ValueError, match="inconsistentes|limite real"):
        _guard_ooxml_zip(bytes(payload))


def test_xlsx_with_xml_entities_is_rejected_not_expanded():
    """Un XLSX con entidades en el XML de la hoja (billion laughs) se rechaza
    sin expandir: el iterparse interno usa defusedxml."""
    from defusedxml.common import EntitiesForbidden

    bomb_sheet = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<!DOCTYPE lolz [<!ENTITY lol "lollollollollollollollollollol">'
        b'<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
        b'<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<dimension ref="A1:B2"/><sheetData><row r="1"><c r="A1"><v>&lol3;</v></c></row></sheetData></worksheet>'
    )
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/worksheets/sheet1.xml", bomb_sheet)
    with pytest.raises((EntitiesForbidden, ValueError)):
        DocumentIngestionService()._parse_native(buf.getvalue(), ".xlsx")
