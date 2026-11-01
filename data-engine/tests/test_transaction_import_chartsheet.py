"""Importación Excel cuando la pestaña activa no es una hoja de datos.

Un .xlsx guardado con la pestaña de un gráfico seleccionada abre con
``wb.active`` apuntando a esa hoja de gráfico (Chartsheet), que no tiene filas.
Antes la importación fallaba con «'Chartsheet' object has no attribute
'iter_rows'» aunque el libro sí tenía las operaciones en otra hoja.
"""

from __future__ import annotations

import io

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.chart import BarChart, Reference  # noqa: E402

from app.services.transaction_import_service import parse_file_rows  # noqa: E402


def _xlsx(active_chart: bool, with_data: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Operaciones"
    if with_data:
        ws.append(["Fecha", "Ticker", "Accion", "Cantidad", "Precio"])
        ws.append(["2026-01-05", "AAPL", "buy", "2", "150.5"])
    else:
        ws.append([None])
    chart_sheet = wb.create_chartsheet("Grafico")
    bar = BarChart()
    bar.add_data(Reference(ws, min_col=4, min_row=1, max_row=2), titles_from_data=True)
    chart_sheet.add_chart(bar)
    if not with_data:
        ws.sheet_state = "hidden"
    if active_chart:
        wb.active = chart_sheet
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_lee_la_hoja_de_datos_aunque_la_activa_sea_un_grafico():
    headers, rows = parse_file_rows(_xlsx(active_chart=True), "operaciones.xlsx")
    assert headers == ["Fecha", "Ticker", "Accion", "Cantidad", "Precio"]
    assert rows == [["2026-01-05", "AAPL", "buy", "2", "150.5"]]


def test_la_hoja_activa_de_datos_sigue_mandando():
    headers, rows = parse_file_rows(_xlsx(active_chart=False), "operaciones.xlsx")
    assert headers[1] == "Ticker"
    assert len(rows) == 1


def test_libro_sin_hojas_de_datos_da_un_mensaje_claro():
    with pytest.raises(ValueError) as excinfo:
        parse_file_rows(_xlsx(active_chart=True, with_data=False), "graficos.xlsx")
    message = str(excinfo.value)
    assert "no contiene filas con datos" in message
    assert "attribute" not in message
