import re
import zipfile
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook

from app.services.document_ingestion_service import DocumentIngestionService, _guard_ooxml_zip


def _xlsx_sparse(dimension=True):
    source = BytesIO()
    book = Workbook()
    book.active['A1'] = 'start'
    book.save(source)
    output = BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as result:
        for name in archive.namelist():
            data = archive.read(name)
            if name == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension[^>]*/>', b'<dimension ref="A1:XFD1048576"/>' if dimension else b'', data)
                data = data.replace(b'</sheetData>', b'<row r="1048576"><c r="XFD1048576" t="inlineStr"><is><t>end</t></is></c></row></sheetData>')
            result.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize('dimension', [True, False])
def test_sparse_xlsx_rejected_before_openpyxl_iteration(monkeypatch, dimension):
    data = _xlsx_sparse(dimension)
    _guard_ooxml_zip(data)
    monkeypatch.setattr('openpyxl.load_workbook', lambda *a, **kw: pytest.fail('must reject before workbook load'))
    with pytest.raises(ValueError, match='dimensiones'):
        DocumentIngestionService()._parse_xlsx(data)


def test_normal_xlsx_preserves_rows():
    book = Workbook()
    book.active.append(['Revenue', 123])
    book.active.append(['FCF', 45])
    data = BytesIO()
    book.save(data)
    parsed = DocumentIngestionService()._parse_xlsx(data.getvalue())
    assert parsed.blocks[0].text == 'Revenue | 123\nFCF | 45'


@pytest.mark.parametrize('paragraph', [False, True])
def test_docx_tables_preserve_figures_and_order(paragraph):
    doc = Document()
    if paragraph:
        doc.add_paragraph('Before')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Revenue'
    table.cell(0, 1).text = '123'
    table.cell(1, 0).text = 'FCF'
    table.cell(1, 1).text = '45'
    if paragraph:
        doc.add_paragraph('After')
    data = BytesIO()
    doc.save(data)
    parsed = DocumentIngestionService()._parse_docx(data.getvalue())
    texts = [block.text for block in parsed.blocks]
    assert texts == (['Before'] if paragraph else []) + ['Revenue | 123\nFCF | 45'] + (['After'] if paragraph else [])
    chunks = DocumentIngestionService()._chunk_blocks(parsed.blocks, 'test', parsed.parser, 'test.docx', None)
    assert chunks and '123' in str(chunks) and '45' in str(chunks)


def test_xlsx_document_budget_applies_across_sheets(monkeypatch):
    book = Workbook()
    for sheet in [book.active, book.create_sheet()]:
        sheet['B2'] = 'x'
    data = BytesIO()
    book.save(data)
    monkeypatch.setattr('app.services.document_ingestion_service.MAX_XLSX_CELLS', 7)
    with pytest.raises(ValueError, match='dimensiones'):
        DocumentIngestionService()._parse_xlsx(data.getvalue())


def test_docx_merged_and_nested_table_cells_preserved_once():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = 'Merged revenue 123'
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'Nested FCF 45'
    data = BytesIO()
    doc.save(data)
    text = DocumentIngestionService()._parse_docx(data.getvalue()).blocks[0].text
    assert text.count('Merged revenue 123') == 1
    assert text.count('Nested FCF 45') == 1
